"""BankSalad workbook and generic ledger CSV import."""

from __future__ import annotations

import csv
import hashlib
import io
from datetime import date, datetime, time, timedelta
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .history import create_portfolio_snapshot
from .models import AssetItem, Debt, Holding, ImportBatch, Transaction


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _number(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = str(value).replace(",", "").replace("원", "").replace("%", "").strip()
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    raw = _text(value).replace(".", "-").replace("/", "-")
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%y-%m-%d"):
        try:
            return datetime.strptime(raw[:19], fmt).date()
        except ValueError:
            continue
    return None


def _time(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%H:%M:%S")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, (int, float)):
        seconds = round(float(value) * 86400) % 86400
        return f"{seconds // 3600:02d}:{(seconds % 3600) // 60:02d}:{seconds % 60:02d}"
    return _text(value)[:20]


def _hash(*parts: Any) -> str:
    raw = "|".join(_text(part).lower() for part in parts)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _find_sheet(workbook, preferred: str, required_headers: set[str]):  # type: ignore[no-untyped-def]
    if preferred in workbook.sheetnames:
        return workbook[preferred]
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 30), values_only=True):
            values = {_text(value) for value in row}
            if required_headers.issubset(values):
                return sheet
    return None


def _header_map(row: tuple[Any, ...]) -> dict[str, int]:
    return {_text(value): index for index, value in enumerate(row) if _text(value)}


def _first_index(headers: dict[str, int], names: tuple[str, ...]) -> int | None:
    for name in names:
        if name in headers:
            return headers[name]
    return None


def _parse_transactions_sheet(sheet) -> list[dict[str, Any]]:  # type: ignore[no-untyped-def]
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next(
        (index for index, row in enumerate(rows[:30]) if "날짜" in {_text(value) for value in row}),
        None,
    )
    if header_index is None:
        return []
    headers = _header_map(rows[header_index])
    date_i = _first_index(headers, ("날짜", "이용일", "결제일"))
    merchant_i = _first_index(headers, ("내용", "가맹점명", "이용가맹점", "상호"))
    amount_i = _first_index(headers, ("금액", "이용금액", "결제금액"))
    if date_i is None or merchant_i is None or amount_i is None:
        return []

    def value(row: tuple[Any, ...], index: int | None) -> Any:
        return row[index] if index is not None and index < len(row) else None

    time_i = _first_index(headers, ("시간",))
    type_i = _first_index(headers, ("타입", "구분"))
    primary_i = _first_index(headers, ("대분류", "카테고리"))
    secondary_i = _first_index(headers, ("소분류",))
    currency_i = _first_index(headers, ("화폐", "통화"))
    payment_i = _first_index(headers, ("결제수단", "카드명"))
    memo_i = _first_index(headers, ("메모", "비고"))
    parsed: list[dict[str, Any]] = []
    for row in rows[header_index + 1 :]:
        transaction_date = _date(value(row, date_i))
        merchant = _text(value(row, merchant_i))
        signed_amount = _number(value(row, amount_i))
        if not transaction_date or not merchant or signed_amount == 0:
            continue
        transaction_type = _text(value(row, type_i)) or ("지출" if signed_amount < 0 else "수입")
        parsed.append(
            {
                "transaction_date": transaction_date,
                "transaction_time": _time(value(row, time_i)),
                "transaction_type": transaction_type,
                "primary_category": _text(value(row, primary_i)) or "미분류",
                "secondary_category": _text(value(row, secondary_i)) or "미분류",
                "merchant": merchant,
                "amount": abs(signed_amount),
                "signed_amount": signed_amount,
                "currency": _text(value(row, currency_i)) or "KRW",
                "payment_method": _text(value(row, payment_i)),
                "memo": _text(value(row, memo_i)),
            }
        )
    return parsed


def _section_rows(sheet, label: str, next_label: str | None = None):  # type: ignore[no-untyped-def]
    rows = list(sheet.iter_rows(values_only=True))
    start = next((index for index, row in enumerate(rows) if any(_text(value) == label for value in row)), None)
    if start is None:
        return []
    end = len(rows)
    if next_label:
        end = next(
            (index for index, row in enumerate(rows[start + 1 :], start + 1) if any(_text(value) == next_label for value in row)),
            end,
        )
    return rows[start:end]


def _parse_financial_status(sheet) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:  # type: ignore[no-untyped-def]
    rows = _section_rows(sheet, "3.재무현황", "4.신용현황")
    assets: list[dict[str, Any]] = []
    debts: list[dict[str, Any]] = []
    asset_category = ""
    debt_category = ""
    skip = {"항목", "총자산", "순자산", "총부채", "자산", "부채"}
    for row in rows:
        padded = tuple(row) + (None,) * 10
        if _text(padded[1]) and _text(padded[1]) not in skip and not _text(padded[1]).startswith("3."):
            asset_category = _text(padded[1])
        asset_name = _text(padded[2])
        asset_value = _number(padded[4])
        if asset_name and asset_value and asset_category not in skip:
            assets.append({"category": asset_category, "provider": "", "name": asset_name, "value": asset_value})

        if _text(padded[5]) and _text(padded[5]) not in skip:
            debt_category = _text(padded[5])
        debt_name = _text(padded[6])
        debt_balance = _number(padded[8])
        if debt_name and debt_balance and debt_category not in skip:
            debts.append(
                {
                    "debt_type": debt_category,
                    "provider": "",
                    "name": debt_name,
                    "principal": debt_balance,
                    "balance": debt_balance,
                    "interest_rate": 0,
                    "opened_on": None,
                    "matures_on": None,
                }
            )
    # BankSalad can export several accounts with the same generic product name
    # (for example, multiple rows named "위탁"). Aggregate them so the stable
    # key stays deterministic across repeated uploads.
    aggregated_assets: dict[tuple[str, str, str], dict[str, Any]] = {}
    for item in assets:
        key = (item["category"], item["provider"], item["name"])
        if key in aggregated_assets:
            aggregated_assets[key]["value"] += item["value"]
        else:
            aggregated_assets[key] = item.copy()
    return list(aggregated_assets.values()), debts


def _parse_holdings(sheet) -> list[dict[str, Any]]:  # type: ignore[no-untyped-def]
    rows = _section_rows(sheet, "5.투자현황", "6.대출현황")
    holdings: list[dict[str, Any]] = []
    for row in rows:
        padded = tuple(row) + (None,) * 10
        asset_type = _text(padded[1])
        broker = _text(padded[2])
        name = _text(padded[3])
        if asset_type in {"", "투자상품종류", "총계"} or not name or name == "보유상품개수":
            continue
        principal = _number(padded[5])
        market_value = _number(padded[6])
        if not principal and not market_value:
            continue
        return_rate = _number(padded[7]) / 100
        holdings.append(
            {
                "asset_type": asset_type,
                "broker": broker,
                "name": name,
                "principal": principal,
                "market_value": market_value,
                "return_rate": return_rate,
            }
        )
    return holdings


def _parse_generic_holdings_sheet(sheet) -> list[dict[str, Any]]:  # type: ignore[no-untyped-def]
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next(
        (index for index, row in enumerate(rows[:30]) if {"종목명", "상품명", "종목", "보유종목"} & {_text(value).replace(" ", "") for value in row}),
        None,
    )
    if header_index is None:
        return []
    headers = {_text(value).replace(" ", ""): index for index, value in enumerate(rows[header_index]) if _text(value)}

    def first(names: tuple[str, ...]) -> int | None:
        return _first_index(headers, names)

    def value(row: tuple[Any, ...], index: int | None) -> Any:
        return row[index] if index is not None and index < len(row) else None

    name_i = first(("종목명", "상품명", "종목", "보유종목"))
    ticker_i = first(("종목코드", "티커", "코드", "Symbol"))
    quantity_i = first(("보유수량", "잔고수량", "수량", "Quantity"))
    average_i = first(("평균단가", "매입단가", "평균매입가", "매수평균가", "AvgPrice"))
    current_i = first(("현재가", "평가단가", "종가", "CurrentPrice"))
    principal_i = first(("투자원금", "매입금액", "원금", "InvestedAmount"))
    valuation_i = first(("평가금액", "평가액", "현재금액", "MarketValue"))
    currency_i = first(("통화", "화폐", "Currency"))
    broker_i = first(("금융사", "증권사", "계좌", "Broker"))
    if name_i is None:
        return []

    holdings: list[dict[str, Any]] = []
    for row in rows[header_index + 1 :]:
        name = _text(value(row, name_i))
        ticker = _text(value(row, ticker_i)).upper()
        quantity = abs(_number(value(row, quantity_i)))
        average = abs(_number(value(row, average_i)))
        current = abs(_number(value(row, current_i)))
        principal = abs(_number(value(row, principal_i))) or quantity * average
        market_value = abs(_number(value(row, valuation_i))) or quantity * current
        if not name or (not principal and not market_value):
            continue
        raw_currency = _text(value(row, currency_i)).upper()
        currency = "USD" if "USD" in raw_currency or (ticker and not ticker.isdigit() and not ticker.endswith((".KS", ".KQ"))) else "KRW"
        holdings.append({
            "asset_type": "주식/ETF",
            "broker": _text(value(row, broker_i)),
            "name": name,
            "ticker": ticker,
            "ticker_source": "user" if ticker else "",
            "currency": currency,
            "principal": principal,
            "market_value": market_value,
            "return_rate": (market_value - principal) / principal if principal else 0,
            "quantity": quantity or None,
            "quantity_source": "upload" if quantity else "",
            "current_price": current or None,
            "price_source": "upload",
        })
    return holdings


def _parse_debt_section(sheet) -> list[dict[str, Any]]:  # type: ignore[no-untyped-def]
    rows = _section_rows(sheet, "6.대출현황")
    debts: list[dict[str, Any]] = []
    for row in rows:
        padded = tuple(row) + (None,) * 10
        debt_type = _text(padded[1])
        provider = _text(padded[2])
        name = _text(padded[3])
        if debt_type in {"", "대출종류", "총계"} or not name or name == "보유 대출 상품":
            continue
        balance = _number(padded[6])
        if not balance:
            continue
        debts.append(
            {
                "debt_type": debt_type,
                "provider": provider,
                "name": name,
                "principal": _number(padded[5]),
                "balance": balance,
                "interest_rate": _number(padded[7]) / 100,
                "opened_on": _date(padded[8]),
                "matures_on": _date(padded[9]),
            }
        )
    return debts


def parse_upload(content: bytes, filename: str) -> dict[str, list[dict[str, Any]]]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        decoded = None
        for encoding in ("utf-8-sig", "cp949", "utf-8"):
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None:
            raise ValueError("CSV 인코딩을 읽을 수 없습니다.")
        rows = list(csv.reader(io.StringIO(decoded)))
        if not rows:
            raise ValueError("비어 있는 CSV입니다.")
        # Reuse the workbook parser by creating a tiny in-memory workbook.
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        for row in rows:
            sheet.append(row)
        holdings = _parse_generic_holdings_sheet(sheet)
        return {"transactions": [] if holdings else _parse_transactions_sheet(sheet), "assets": [], "holdings": holdings, "debts": []}

    if not lower.endswith((".xlsx", ".xlsm")):
        raise ValueError("현재 서버 업로드는 .xlsx, .xlsm, .csv 파일을 지원합니다.")
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ledger = _find_sheet(workbook, "가계부 내역", {"날짜", "금액"})
    status = _find_sheet(workbook, "뱅샐현황", {"5.투자현황"})
    transactions = _parse_transactions_sheet(ledger) if ledger else []
    assets, status_debts = _parse_financial_status(status) if status else ([], [])
    holdings = _parse_holdings(status) if status else []
    detailed_debts = _parse_debt_section(status) if status else []
    if not status:
        generic_sheet = workbook[workbook.sheetnames[0]]
        holdings = _parse_generic_holdings_sheet(generic_sheet)
        if holdings:
            transactions = []
    return {
        "transactions": transactions,
        "assets": assets,
        "holdings": holdings,
        "debts": detailed_debts or status_debts,
    }


def _deactivate(db: Session, model, owner: str) -> None:  # type: ignore[no-untyped-def]
    for item in db.scalars(select(model).where(model.owner == owner, model.is_active.is_(True))):
        item.is_active = False


def import_upload(db: Session, content: bytes, filename: str, owner: str) -> dict[str, Any]:
    parsed = parse_upload(content, filename)
    checksum = hashlib.sha256(content).hexdigest()
    batch = ImportBatch(filename=filename, checksum=checksum, owner=owner)
    db.add(batch)
    db.flush()

    existing_keys = set(db.scalars(select(Transaction.source_key)))
    new_transactions = 0
    for item in parsed["transactions"]:
        source_key = _hash(
            owner,
            item["transaction_date"],
            item["transaction_time"],
            item["transaction_type"],
            item["merchant"],
            item["signed_amount"],
            item["payment_method"],
        )
        if source_key in existing_keys:
            continue
        db.add(Transaction(source_key=source_key, import_batch_id=batch.id, owner=owner, **item))
        existing_keys.add(source_key)
        new_transactions += 1

    if parsed["assets"]:
        _deactivate(db, AssetItem, owner)
        existing_assets = {
            item.source_key: item for item in db.scalars(select(AssetItem).where(AssetItem.owner == owner))
        }
        for item in parsed["assets"]:
            source_key = _hash(item["category"], item["provider"], item["name"])
            record = existing_assets.get(source_key)
            if record:
                for key, value in item.items():
                    setattr(record, key, value)
                record.is_active = True
            else:
                db.add(AssetItem(owner=owner, source_key=source_key, is_active=True, **item))

    if parsed["holdings"]:
        _deactivate(db, Holding, owner)
        existing_holdings = {
            item.source_key: item for item in db.scalars(select(Holding).where(Holding.owner == owner))
        }
        for item in parsed["holdings"]:
            source_key = _hash(item["broker"], item["name"])
            record = existing_holdings.get(source_key)
            if record:
                preserved = {
                    "ticker": record.ticker,
                    "ticker_source": record.ticker_source,
                    "quantity": record.quantity,
                    "quantity_source": record.quantity_source,
                    "target_price": record.target_price,
                    "target_alert_enabled": record.target_alert_enabled,
                    "target_alert_sent_at": record.target_alert_sent_at,
                }
                for key, value in item.items():
                    setattr(record, key, value)
                for key, value in preserved.items():
                    if key not in item:
                        setattr(record, key, value)
                record.price_source = "bank_salad"
                record.is_active = True
            else:
                db.add(Holding(owner=owner, source_key=source_key, is_active=True, **item))

    if parsed["debts"]:
        _deactivate(db, Debt, owner)
        existing_debts = {
            item.source_key: item for item in db.scalars(select(Debt).where(Debt.owner == owner))
        }
        for item in parsed["debts"]:
            source_key = _hash(item["provider"], item["name"])
            record = existing_debts.get(source_key)
            if record:
                for key, value in item.items():
                    setattr(record, key, value)
                record.is_active = True
            else:
                db.add(Debt(owner=owner, source_key=source_key, is_active=True, **item))

    db.flush()
    batch.transaction_count = len(parsed["transactions"])
    batch.new_transaction_count = new_transactions
    batch.holding_count = len(parsed["holdings"])
    batch.asset_count = len(parsed["assets"])
    snapshot = create_portfolio_snapshot(db, source="bank_salad_import", import_batch_id=batch.id)
    db.commit()
    return {
        "batch_id": batch.id,
        "snapshot_id": snapshot.id,
        "transactions_read": len(parsed["transactions"]),
        "transactions_added": new_transactions,
        "holdings_updated": len(parsed["holdings"]),
        "assets_updated": len(parsed["assets"]),
        "debts_updated": len(parsed["debts"]),
    }
